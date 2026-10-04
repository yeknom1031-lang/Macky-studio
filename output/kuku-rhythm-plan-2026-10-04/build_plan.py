import json
from pathlib import Path
from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
data = json.loads((ROOT / 'plan-content.json').read_text())
doc = Document()
s = doc.sections[0]
s.page_width = Mm(210)
s.page_height = Mm(297)
s.top_margin = Mm(17)
s.bottom_margin = Mm(17)
s.left_margin = Mm(19)
s.right_margin = Mm(19)
s.footer_distance = Mm(8)
for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Caption']:
    style = doc.styles[name]
    style.font.name = 'YuGothic'
    fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn('w:eastAsia'), 'YuGothic')
    for attr in list(fonts.attrib):
        if 'theme' in attr.lower():
            del fonts.attrib[attr]
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_after = Pt(6)
    style.paragraph_format.line_spacing = Pt(14.5)
    style.paragraph_format.space_before = Pt(0)
    snap = OxmlElement('w:snapToGrid')
    snap.set(qn('w:val'), '0')
    style.element.get_or_add_pPr().append(snap)
doc.styles['Normal'].font.size = Pt(10)
doc.styles['Title'].font.size = Pt(24)
doc.styles['Title'].font.bold = True
doc.styles['Title'].paragraph_format.line_spacing = Pt(31)
doc.styles['Heading 1'].font.size = Pt(18)
doc.styles['Heading 1'].font.bold = True
doc.styles['Heading 1'].paragraph_format.line_spacing = Pt(24)
doc.styles['Heading 1'].paragraph_format.space_after = Pt(11)
doc.styles['Heading 2'].font.size = Pt(11.5)
doc.styles['Heading 2'].font.bold = True
doc.styles['Heading 2'].paragraph_format.space_before = Pt(9)
doc.styles['Heading 2'].paragraph_format.space_after = Pt(4)
doc.styles['Subtitle'].font.size = Pt(12)
doc.styles['Caption'].font.size = Pt(8)
doc.styles['Caption'].font.italic = False
doc.styles['Caption'].paragraph_format.line_spacing = Pt(11)
doc.styles['Caption'].font.bold = False
for style in doc.styles:
    for border in style.element.xpath('.//w:pBdr'):
        border.getparent().remove(border)
doc.core_properties.title = data['title']
doc.core_properties.subject = '九九学習リズムゲームの企画と制作計画'
doc.core_properties.author = 'Macky Studio'

footer = s.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r = footer.add_run()
r.font.size = Pt(8)
f = OxmlElement('w:fldSimple')
f.set(qn('w:instr'), 'PAGE')
r._r.addnext(f)

def p(text, style=None):
    paragraph = doc.add_paragraph(text, style)
    paragraph.paragraph_format.widow_control = True
    return paragraph

def table(block):
    rows = [block['headers']] + block['rows']
    t = doc.add_table(rows=0, cols=len(rows[0]))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    widths = block.get('widths', [172 / len(rows[0])] * len(rows[0]))
    for col, width in zip(t.columns, widths):
        col.width = Mm(width)
    borders = OxmlElement('w:tblBorders')
    for edge in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        e = OxmlElement('w:' + edge)
        e.set(qn('w:val'), 'single')
        e.set(qn('w:sz'), '4')
        e.set(qn('w:color'), 'D9D9D9')
        borders.append(e)
    t._tbl.tblPr.append(borders)
    for i, values in enumerate(rows):
        row = t.add_row()
        props = row._tr.get_or_add_trPr()
        props.append(OxmlElement('w:cantSplit'))
        if i == 0:
            props.append(OxmlElement('w:tblHeader'))
        for j, (c, value) in enumerate(zip(row.cells, values)):
            c.width = Mm(widths[j])
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr = c._tc.get_or_add_tcPr()
            margins = OxmlElement('w:tcMar')
            for edge in ['top', 'left', 'bottom', 'right']:
                e = OxmlElement('w:' + edge)
                e.set(qn('w:w'), '85')
                e.set(qn('w:type'), 'dxa')
                margins.append(e)
            tcpr.append(margins)
            shade = OxmlElement('w:shd')
            shade.set(qn('w:fill'), 'DCEAF1' if i == 0 else ('F6F8FA' if i % 2 == 0 else 'FFFFFF'))
            tcpr.append(shade)
            paragraph = c.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = Pt(13)
            if j in block.get('center', []):
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = paragraph.add_run(str(value))
            run.font.size = Pt(block.get('size', 9))
            run.font.bold = i == 0
    p('').paragraph_format.space_after = Pt(0)

for idx, page in enumerate(data['pages']):
    if idx:
        doc.add_page_break()
    p(page['title'], 'Title' if idx == 0 else 'Heading 1')
    for block in page['blocks']:
        kind = block['type']
        if kind == 'p':
            p(block['text'])
        elif kind == 'h':
            p(block['text'], 'Heading 2')
        elif kind == 'sub':
            p(block['text'], 'Subtitle')
        elif kind == 'table':
            table(block)
        elif kind == 'image':
            para = doc.add_paragraph()
            para.paragraph_format.line_spacing = 1
            run = para.add_run()
            shape = run.add_picture(str(ROOT / block['path']), width=Mm(172))
            shape._inline.docPr.set('descr', block['alt'])
            p(block['caption'], 'Caption')
        elif kind == 'list':
            for text in block['items']:
                paragraph = p('・' + text)
                paragraph.paragraph_format.space_after = Pt(4)
        elif kind == 'sources':
            for item in block['items']:
                paragraph = p(item['title'])
                paragraph.paragraph_format.space_after = Pt(2)
                rel = doc.part.relate_to(item['url'], 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)
                link = OxmlElement('w:hyperlink')
                link.set(qn('r:id'), rel)
                r = OxmlElement('w:r')
                rp = OxmlElement('w:rPr')
                color = OxmlElement('w:color'); color.set(qn('w:val'), '155D84'); rp.append(color)
                size = OxmlElement('w:sz'); size.set(qn('w:val'), '17'); rp.append(size)
                r.append(rp)
                text = OxmlElement('w:t'); text.text = item['url']; r.append(text)
                link.append(r)
                p('')._p.append(link)
                if item.get('note'):
                    p(item['note'], 'Caption')

out = ROOT / '九九リズムゲーム企画書.docx'
doc.save(out)
print(out)
